"""Pipeline orchestrator: the only place where the four modules meet.

    Submission ──► [1 Shield] ──► Artifacts (pseudonymised)        ledger: ingestion
                                     │
                                     ▼
                   [2 Funnel L1] ──► Signals + credentials + density   (local, free)
                                     │
                   [2 Funnel L2] ──► top-N% only, token & $ budget      ledger: escalation
                                     │
                   [3 Translator] ─► Skill graph (3 axes, evidence)
                                     │
                   [4 Dashboard] ──► Report (score, proofs, guide)      ledger: score
                                     │
                   Human reviewer ─► Decision                           ledger: human_decision

Modules never import each other sideways; data crosses boundaries only as models from ``models``.
"""

from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import secrets
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

from .config import ENGINE_VERSION, Settings
from .dashboard.interview import build_interview_guide
from .dashboard.scoring import score
from .funnel.budget import BudgetGuard, Price, estimate_tokens, factual_density, plan_escalation
from .funnel.crossrepo import analyse_repository, detect_links
from .funnel.documents import analyze_text, extract_credentials
from .funnel.github import FetchedRepo, fetch_repositories
from .funnel.llm import LLMError, LLMProvider, build_provider
from .funnel.repo import RepoFetchError, RepoSnapshot, analyze_repo
from .i18n import localise
from .models import (
    Artifact,
    ArtifactKind,
    ArtifactStatus,
    AuditInfo,
    Candidate,
    CandidateSummary,
    ConsentRecord,
    CredentialItem,
    DashboardReport,
    EscalationRecord,
    HumanDecision,
    InterviewQuestion,
    JobProfile,
    L1Report,
    RedactionReport,
    Signal,
    SkillGraph,
    utcnow,
)
from .pilot.scoring import PILOT_BLEND
from .shield.injection import hidden_count, screen, strip_hidden
from .shield.keys import resolve_seal_key, resolve_vault_key
from .shield.ledger import AuditLedger, ChainVerification, LedgerEntry
from .shield.ner import build_ner
from .shield.pii import TextPseudonymizer
from .shield.vault import PseudonymVault
from .shield.vision import DetectorUnavailable, RegionDetector, build_detector, redact_image
from .store import Store
from .translator.catalog import CATALOG
from .translator.heuristic import build_skill_graph
from .translator.llm_eval import merge_llm_output
from .translator.prompts import LLM_OUTPUT_SCHEMA, build_user_message, system_prompt

log = logging.getLogger(__name__)

NOTICE = {
    "fr": ("Aide à la décision uniquement. Ce score ne rejette personne : il classe des preuves. Toute décision "
           "est prise par une personne identifiée, qui la motive, et chaque étape est consignée dans un registre "
           "d'audit infalsifiable (AI Act, art. 12-14 ; RGPD, art. 22)."),
    "en": ("Decision support only. This score rejects no one: it ranks evidence. Every decision is made and "
           "justified by an identified person, and every step is recorded in a tamper-evident audit ledger "
           "(EU AI Act Art. 12-14; GDPR Art. 22)."),
}


# ================================================================================================
# Submission inputs
# ================================================================================================


class TextDocument(BaseModel):
    name: str = Field(..., max_length=255)
    content: str = Field(..., max_length=400_000)
    kind: ArtifactKind = ArtifactKind.document


class RepositoryInput(BaseModel):
    url: str | None = None
    # Integrations may send a pre-computed tree (paths only) and optional key files instead of a URL.
    paths: list[str] | None = Field(None, max_length=20_000)
    files: dict[str, str] | None = None


class PortfolioItem(BaseModel):
    title: str = Field("", max_length=200)
    description: str = Field(..., max_length=8000)


class Submission(BaseModel):
    consent: bool
    # Required: an exact match on the declared name is the only defence that works whatever the CV layout.
    # Measured on unseen layouts, rules alone caught 1.4% of undeclared names (docs/MEASUREMENTS.md).
    identity_name: str = Field(..., min_length=2, max_length=200,
                               description="Used only to mask the name everywhere, then stored encrypted.")
    identity_email: str = Field("", max_length=200)
    retention_days: int = Field(180, ge=1, le=730)
    documents: list[TextDocument] = Field(default_factory=list, max_length=60)
    repositories: list[RepositoryInput] = Field(default_factory=list, max_length=10)
    portfolio: list[PortfolioItem] = Field(default_factory=list, max_length=30)


@dataclass
class ImageInput:
    name: str
    data: bytes
    caption: str = ""


@dataclass
class EvaluationRun:
    job_id: str
    evaluated: int
    escalated: list[str] = field(default_factory=list)
    skipped: dict[str, str] = field(default_factory=dict)
    spent_usd: float = 0.0


class NotFound(LookupError):
    pass


class PolicyError(ValueError):
    pass


def extract_text(name: str, data: bytes) -> str:
    """Plain text from an uploaded document (PDF, Word .docx, Markdown, text)."""
    lower = name.lower()
    if lower.endswith(".pdf") or data[:5] == b"%PDF-":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages[:60]]
        except (PdfReadError, ValueError) as exc:
            raise PolicyError(f"{name}: unreadable PDF ({exc})") from exc
        # Layout mode keeps multi-column CVs readable when the default extraction collapses them.
        if sum(len(p) for p in pages) < 200:
            pages = [page.extract_text(extraction_mode="layout") or "" for page in reader.pages[:60]]
        return "\n".join(pages)
    if lower.endswith(".docx") or (data[:2] == b"PK" and b"word/document.xml" in data[:4000]):
        import zipfile

        from docx import Document

        try:
            doc = Document(io.BytesIO(data))
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise PolicyError(f"{name}: unreadable Word document ({exc})") from exc
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many CV templates put the whole layout in tables
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    if cell.text not in cells:  # merged cells repeat their text
                        cells.append(cell.text)
                lines.append(" | ".join(cells))
        for section in doc.sections:  # contact details often live in the header
            lines = [p.text for p in section.header.paragraphs] + lines
        return "\n".join(lines)
    if lower.endswith((".doc", ".odt", ".rtf", ".pages")):
        raise PolicyError(f"{name}: unsupported format, please send PDF, DOCX, Markdown or plain text")
    return data.decode("utf-8", errors="replace")


# ================================================================================================
# Engine
# ================================================================================================


class Engine:
    def __init__(
        self,
        settings: Settings,
        store: Store | None = None,
        detector: RegionDetector | None = None,
        provider: LLMProvider | bool | None = True,
    ) -> None:
        self.settings = settings
        settings.data_dir.mkdir(parents=True, exist_ok=True)
        self.store = store or Store(settings.db_path)
        self.vault_key = resolve_vault_key(settings.vault_key, settings.data_dir)
        self.vault = PseudonymVault(self.store, self.vault_key)
        self.ledger = AuditLedger(self.store, resolve_seal_key(settings.ledger_seal_key, settings.data_dir))
        self.detector = detector or build_detector(settings.vision_detector, settings.ollama_url,
                                                   settings.vision_model)
        self.provider: LLMProvider | None = build_provider(settings) if provider is True else (provider or None)
        self.budget = BudgetGuard(self.store, Price(settings.llm_price_input_per_mtok,
                                                    settings.llm_price_output_per_mtok))
        self.ner = build_ner(settings.ner)
        self.media_dir = settings.data_dir / "media"
        self.media_dir.mkdir(parents=True, exist_ok=True)

    # -------------------------------------------------------------------------------------- jobs

    def create_job(self, job: JobProfile, actor: str = "recruiter") -> JobProfile:
        _validate_skills(job)
        job = job.model_copy(update={"id": f"JOB-{secrets.token_hex(3).upper()}", "version": 1,
                                     "created_at": utcnow(), "updated_at": utcnow()})
        self.store.put("jobs", job.id, job)
        self.ledger.append("job_config", {"version": 1, "config": job.model_dump(mode="json")}, actor=actor,
                           job_id=job.id)
        return job

    def update_job(self, job_id: str, job: JobProfile, actor: str = "recruiter") -> JobProfile:
        current = self.get_job(job_id)
        _validate_skills(job)
        job = job.model_copy(update={"id": job_id, "version": current.version + 1,
                                     "created_at": current.created_at, "updated_at": utcnow()})
        self.store.put("jobs", job_id, job)
        self.ledger.append("job_config", {"version": job.version, "config": job.model_dump(mode="json")},
                           actor=actor, job_id=job_id)
        return job

    def get_job(self, job_id: str) -> JobProfile:
        job = self.store.get("jobs", job_id, JobProfile)
        if job is None:
            raise NotFound(f"job {job_id} not found")
        return job

    def list_jobs(self) -> list[JobProfile]:
        return self.store.list("jobs", JobProfile)

    # -------------------------------------------------------------------------------------- module 1

    def ingest(self, job_id: str, sub: Submission, images: list[ImageInput] | None = None) -> Candidate:
        job = self.get_job(job_id)
        if not sub.consent:
            raise PolicyError("explicit consent is required before any processing (GDPR Art. 6 and 13)")
        images = images or []
        if not (sub.documents or sub.repositories or sub.portfolio or images):
            raise PolicyError("a submission needs at least one document, repository, portfolio item or image")
        ref = f"CAND-{secrets.token_hex(3).upper()}"
        known = [v for v in (sub.identity_name, sub.identity_email) if v]
        for kind, value in (("IDENTITY_NAME", sub.identity_name), ("IDENTITY_EMAIL", sub.identity_email)):
            if value:
                self.vault.token(kind, value, ref)

        text_shield = TextPseudonymizer(self.vault, mask_school_names=job.privacy.mask_school_names,
                                        neutralise_gendered_terms=job.privacy.neutralize_gendered_terms,
                                        ner=self.ner)
        path_shield = TextPseudonymizer(self.vault, mask_school_names=False, neutralise_gendered_terms=False)
        artifacts: list[Artifact] = []
        counters: dict[str, int] = {}

        def next_id(kind: ArtifactKind, label_fr: str) -> tuple[str, str]:
            counters[kind] = counters.get(kind, 0) + 1
            n = counters[kind]
            label = label_fr if kind in (ArtifactKind.cv, ArtifactKind.linkedin) and n == 1 else f"{label_fr} {n}"
            return f"A{len(artifacts) + 1}", label

        def shield_text(raw: str, *, header: bool) -> tuple[str, RedactionReport]:
            cleaned = strip_hidden(raw)
            result = text_shield.run(cleaned, ref, known, header_name=header)
            text = result.text
            for rx, repo_label in name_rx:  # "Built agent-platform" in a CV would lead to the GitHub account
                text = rx.sub(repo_label, text)
            # Excerpts come from the pseudonymised text (no PII); invisible characters are counted on the raw one.
            hits = screen(text, hidden=hidden_count(raw))
            return text, RedactionReport(
                pii_replaced=dict(result.counts), gendered_terms_neutralised=result.gendered_terms,
                injection_suspected=bool(hits), injection_excerpts=hits, detector="rules",
            )

        # Repositories first: the GitHub user name they reveal must be masked in every other document too.
        fetched: list[FetchedRepo] = []
        try:
            for repo in sub.repositories:
                if repo.paths is not None:
                    fetched.append(FetchedRepo(url=repo.url or "", owner="", name="",
                                               snapshot=RepoSnapshot(list(repo.paths), dict(repo.files or {}))))
            urls = [r.url for r in sub.repositories if r.paths is None and r.url]
            if urls:
                fetched += fetch_repositories(urls, self.settings.github_token, self.settings.max_repos_per_submission,
                                              sources=self.settings.ownership_source_files)
            elif not fetched and sub.repositories:
                raise PolicyError("a repository needs a url or a list of paths")
        except RepoFetchError as exc:
            raise PolicyError(str(exc)) from exc
        for owner in sorted({f.owner for f in fetched if f.owner}):
            self.vault.token("GITHUB_USER", owner, ref)
            known.append(owner)
        repo_labels = {f.url or str(i): f"repo-{i}" for i, f in enumerate(fetched, start=1)}
        # Repository names can be searched online and would break blind review: excerpts say "repo-3" instead.
        name_rx = [(re.compile(r"(?i)(?<![\w-])" + re.escape(f.name) + r"(?![\w-])"), repo_labels[f.url])
                   for f in sorted(fetched, key=lambda f: -len(f.name)) if f.name and len(f.name) >= 3]

        def blind(text: str) -> str:
            text = text_shield.run(strip_hidden(text), ref, known, header_name=False).text
            for rx, label in name_rx:
                text = rx.sub(label, text)
            return text

        for doc in sub.documents:
            labels = {ArtifactKind.cv: "CV", ArtifactKind.portfolio_note: "Note", ArtifactKind.document: "Document",
                      ArtifactKind.linkedin: "LinkedIn", ArtifactKind.degree: "Diplôme",
                      ArtifactKind.certification: "Certification"}
            aid, label = next_id(doc.kind, labels.get(doc.kind, "Document"))
            self.vault.token("FILENAME", doc.name, ref)
            text, report = shield_text(doc.content, header=doc.kind in (ArtifactKind.cv, ArtifactKind.linkedin))
            artifacts.append(Artifact(id=aid, candidate_ref=ref, kind=doc.kind, label=label, text=text,
                                      content_sha256=_sha(text), redaction=report, media_type="text/plain"))

        for item in sub.portfolio:
            aid, label = next_id(ArtifactKind.portfolio_note, "Projet")
            text, report = shield_text(f"{item.title}\n{item.description}".strip(), header=False)
            artifacts.append(Artifact(id=aid, candidate_ref=ref, kind=ArtifactKind.portfolio_note, label=label,
                                      text=text, content_sha256=_sha(text), redaction=report,
                                      media_type="text/plain"))

        links = detect_links([f for f in fetched if f.owner])
        for i, fetched_repo in enumerate(fetched, start=1):
            aid = f"A{len(artifacts) + 1}"
            label = repo_labels[fetched_repo.url or str(i)]
            if fetched_repo.url:
                self.vault.token("REPO_URL", fetched_repo.url, ref)
            snapshot = fetched_repo.snapshot
            paths = path_shield.run("\n".join(snapshot.paths), ref, known, header_name=False).text.splitlines()
            files: dict[str, str] = {}
            hits: list[str] = []
            pii: dict[str, int] = {}
            for path, content in snapshot.files.items():
                shielded = text_shield.run(strip_hidden(content), ref, known, header_name=False)
                text = shielded.text
                for rx, other in name_rx:
                    text = rx.sub(other, text)
                files[path_shield.run(path, ref, known, header_name=False).text] = text
                hits += screen(text, hidden=hidden_count(content))
                for k, v in shielded.counts.items():
                    pii[k] = pii.get(k, 0) + v
            sources: dict[str, str] = {}
            for path, content in list(fetched_repo.sources.items())[:self.settings.ownership_source_files]:
                shielded = text_shield.run(strip_hidden(content), ref, known, header_name=False)
                text = shielded.text
                for rx, other in name_rx:
                    text = rx.sub(other, text)
                sources[path_shield.run(path, ref, known, header_name=False).text] = text
            integration: dict[str, Any] = {}
            if fetched_repo.owner:
                analysis = analyse_repository(fetched_repo)
                integration = {
                    "links": [{"target": repo_labels[link.target], "kind": link.kind, "file": link.file,
                               "line": link.line, "excerpt": blind(link.excerpt)}
                              for link in links if link.source == fetched_repo.url],
                    "linked_from": sorted({repo_labels[link.source] for link in links
                                           if link.target == fetched_repo.url and link.kind != "documentation"}),
                    "tools": analysis.tools, "workflow_files": analysis.workflow_files,
                    "job_dependencies": analysis.job_dependencies, "services": analysis.services,
                    "compose_file": analysis.compose_file,
                    "stack": {skill: sorted({d for d, _ in deps}) for skill, deps in analysis.stack.items()},
                    "stack_files": {skill: sorted({f for _, f in deps}) for skill, deps in analysis.stack.items()},
                }
            report = RedactionReport(pii_replaced=pii, injection_suspected=bool(hits), injection_excerpts=hits[:5],
                                     detector="rules", notes=[f"{len(paths)} paths read, {len(files)} key files kept"])
            artifacts.append(Artifact(id=aid, candidate_ref=ref, kind=ArtifactKind.repository, label=label,
                                      repo_paths=paths, repo_files=files, source_files=sources,
                                      integration=integration,
                                      content_sha256=_sha("\n".join(paths)), redaction=report,
                                      media_type="application/x-git-tree"))

        for image in images:
            aid, label = next_id(ArtifactKind.image, "Image")
            self.vault.token("FILENAME", image.name, ref)
            caption_text, caption_report = shield_text(image.caption, header=False) if image.caption else ("", None)
            try:
                result = redact_image(image.data, self.detector)
            except DetectorUnavailable as exc:
                if job.privacy.quarantine_images_without_detector:
                    artifacts.append(Artifact(
                        id=aid, candidate_ref=ref, kind=ArtifactKind.image, label=label,
                        status=ArtifactStatus.quarantined, content_sha256=_sha(image.data.hex()),
                        redaction=RedactionReport(detector="unavailable", notes=[
                            f"quarantined: {exc}. No face/logo detector could run, so the image was not analysed."]),
                    ))
                    continue
                result = None
            except OSError as exc:
                raise PolicyError(f"unreadable image {label}: {exc}") from exc
            # The local VLM's description of the (already masked) image, pseudonymised like any text.
            vlm_caption = ""
            if result and result.caption:
                vlm_caption = text_shield.run(result.caption, ref, known, header_name=False).text
            text = "\n".join(filter(None, [vlm_caption, caption_text]))
            report = RedactionReport(
                visual_regions_masked=len(result.regions) if result else 0,
                metadata_removed=result.metadata_removed if result else [],
                detector=result.detector if result else "none (policy allows unverified images)",
                pii_replaced=caption_report.pii_replaced if caption_report else {},
                injection_suspected=bool(caption_report and caption_report.injection_suspected),
                injection_excerpts=caption_report.injection_excerpts if caption_report else [],
            )
            media_type = result.media_type if result else "image/png"
            if result:
                ext = media_type.split("/")[-1]
                (self.media_dir / f"{ref}-{aid}.{ext}").write_bytes(result.image_bytes)
            artifacts.append(Artifact(id=aid, candidate_ref=ref, kind=ArtifactKind.image, label=label, text=text,
                                      content_sha256=_sha(result.image_bytes.hex() if result else text),
                                      redaction=report, media_type=media_type))

        consent = ConsentRecord(given=True, retention_days=sub.retention_days)
        candidate = Candidate(ref=ref, job_id=job_id, consent=consent,
                              artifact_ids=[a.id for a in artifacts])
        for art in artifacts:
            self.store.put("artifacts", f"{ref}:{art.id}", art, parent=ref)
        self.store.put("candidates", ref, candidate, parent=job_id)
        self.ledger.append("ingestion", {
            "consent": candidate.consent.model_dump(mode="json"),
            "artifacts": [{"id": a.id, "label": a.label, "kind": a.kind, "status": a.status,
                           "sha256": a.content_sha256, "redaction": a.redaction.model_dump(mode="json")}
                          for a in artifacts],
            "engine_version": ENGINE_VERSION,
        }, job_id=job_id, candidate_ref=ref)
        return candidate

    # -------------------------------------------------------------------------------------- module 2-4

    def artifacts(self, ref: str) -> list[Artifact]:
        return self.store.list("artifacts", Artifact, parent=ref)

    def level1(self, ref: str) -> tuple[L1Report, list[Artifact]]:
        artifacts = self.artifacts(ref)
        signals: list[Signal] = []
        creds: list[CredentialItem] = []
        for art in artifacts:
            if art.status != ArtifactStatus.ready:
                continue
            if art.kind in (ArtifactKind.degree, ArtifactKind.certification):
                creds.extend(extract_credentials(art))  # proof of a credential, not of a skill
                continue
            found = analyze_repo(art) if art.kind == ArtifactKind.repository else analyze_text(art)
            signals.extend(found)
            if art.kind in (ArtifactKind.cv, ArtifactKind.document, ArtifactKind.linkedin):
                creds.extend(extract_credentials(art))
        creds = _merge_credentials(creds)
        text_volume = sum(len(s.evidence.excerpt) for s in signals) + sum(
            len(c) for a in artifacts for c in a.repo_files.values())
        report = L1Report(
            candidate_ref=ref, signals=signals, credentials=creds, density=factual_density(signals),
            token_estimate=estimate_tokens("x" * text_volume),
            quarantined_artifacts=[a.label for a in artifacts if a.status == ArtifactStatus.quarantined],
            injection_flags=[a.label for a in artifacts if a.redaction.injection_suspected],
        )
        return report, artifacts

    def evaluate(self, job_id: str, actor: str = "system") -> EvaluationRun:
        job = self.get_job(job_id)
        candidates = self.store.list("candidates", Candidate, parent=job_id)
        l1: dict[str, tuple[L1Report, list[Artifact]]] = {c.ref: self.level1(c.ref) for c in candidates}
        plan = plan_escalation({ref: r.density for ref, (r, _) in l1.items()}, job.funnel, self.provider is not None)
        run = EvaluationRun(job_id=job_id, evaluated=len(candidates), skipped=dict(plan.skipped))

        for cand in candidates:
            report1, artifacts = l1[cand.ref]
            for ev in [s.evidence for s in report1.signals] + [c.evidence for c in report1.credentials]:
                ev.locator = localise(ev.locator, job.locale)
            self.store.put("l1", cand.ref, report1, parent=job_id)
            graph = build_skill_graph(cand.ref, report1.signals, job.locale)
            escalation = EscalationRecord(reason=plan.skipped.get(cand.ref, ""))
            llm_questions: list[InterviewQuestion] = []
            warnings: list[str] = []
            if cand.ref in plan.selected:
                graph, llm_questions, escalation, notes = self._escalate(job, cand.ref, report1, artifacts, graph,
                                                                         plan.selected[cand.ref])
                warnings += notes
                if escalation.escalated:
                    run.escalated.append(cand.ref)
                    run.spent_usd += escalation.cost_usd
                else:
                    run.skipped[cand.ref] = escalation.reason
            escalation.reason = localise(escalation.reason, job.locale)
            self.store.put("graphs", cand.ref, graph, parent=job_id)
            self._build_report(job, cand, report1, graph, escalation, llm_questions, warnings, actor)
        return run

    def _escalate(
        self, job: JobProfile, ref: str, l1: L1Report, artifacts: list[Artifact], graph: SkillGraph, reason: str,
    ) -> tuple[SkillGraph, list[InterviewQuestion], EscalationRecord, list[str]]:
        assert self.provider is not None
        safe = {a.id: a for a in artifacts if a.status == ArtifactStatus.ready and not a.redaction.injection_suspected}
        signals = [s for s in l1.signals if s.artifact_id in safe]
        message, included = build_user_message(signals, safe, locale=job.locale,
                                               max_input_tokens=job.funnel.max_input_tokens_per_candidate)
        tokens = estimate_tokens(system_prompt()) + estimate_tokens(message)
        ok, why, _worst = self.budget.authorize(job.id, job.funnel, tokens)
        record = EscalationRecord(provider=self.provider.name, model=self.provider.model)
        if not ok:
            record.reason = f"escalation refused by budget guard: {why}"
            self.ledger.append("escalation", {"selected_because": reason, "authorized": False, "reason": why,
                                              "estimated_input_tokens": tokens}, job_id=job.id, candidate_ref=ref)
            return graph, [], record, []
        try:
            result = self.provider.complete_json(system_prompt(), message, LLM_OUTPUT_SCHEMA,
                                                 job.funnel.max_output_tokens_per_candidate)
        except LLMError as exc:
            record.reason = f"escalation failed, Level-1 assessment kept: {exc}"
            self.ledger.append("escalation", {"selected_because": reason, "authorized": True, "error": str(exc)},
                               job_id=job.id, candidate_ref=ref)
            return graph, [], record, []
        self.budget.record(job.id, result.input_tokens, result.output_tokens, result.cost_usd)
        merged = merge_llm_output(result.data, graph, l1.signals, safe, included)
        record = EscalationRecord(escalated=True, reason=reason, provider=self.provider.name, model=result.model,
                                  input_tokens=result.input_tokens, output_tokens=result.output_tokens,
                                  cost_usd=result.cost_usd)
        notes = []
        if merged.injection_suspected:
            notes.append("Le modèle d'approfondissement a repéré de possibles instructions cachées dans les pièces."
                         if job.locale == "fr" else
                         "The escalation model flagged possible instructions hidden in the candidate's material.")
        self.ledger.append("escalation", {
            "selected_because": reason, "authorized": True, "provider": record.provider, "model": record.model,
            "input_tokens": record.input_tokens, "output_tokens": record.output_tokens, "cost_usd": record.cost_usd,
            "evidence_ids_sent": sorted(included), "rejected_outputs": merged.rejected,
            "injection_suspected": merged.injection_suspected,
        }, job_id=job.id, candidate_ref=ref)
        return merged.graph, merged.questions, record, notes

    def _build_report(
        self, job: JobProfile, cand: Candidate, l1: L1Report, graph: SkillGraph, escalation: EscalationRecord,
        llm_questions: list[InterviewQuestion], warnings: list[str], actor: str,
    ) -> DashboardReport:
        compat, skills_pct, creds, criteria, gaps, confidence, band = score(job, graph, l1.credentials)
        guide = list(llm_questions)
        if len(guide) < 3:
            claims = [s for s in l1.signals if s.claim_only]
            for q in build_interview_guide(job, graph, claims=claims):
                if len(guide) >= 3:
                    break
                if all(q.skill_id != g.skill_id for g in guide):
                    guide.append(q)
        fr = job.locale == "fr"
        if l1.quarantined_artifacts:
            warnings.append(("Pièces mises en quarantaine (aucun détecteur de visages disponible) : "
                             if fr else "Quarantined items (no face detector available): ")
                            + ", ".join(l1.quarantined_artifacts))
        if l1.injection_flags:
            warnings.append(("Instructions suspectes détectées dans : " if fr else "Suspicious instructions found in: ")
                            + ", ".join(l1.injection_flags)
                            + (" — exclues de l'analyse par IA, à examiner par une personne."
                               if fr else " — excluded from AI analysis, to be reviewed by a person."))
        crit_ids = {c.skill_id for c in job.criteria}
        validated = []
        for a in sorted(graph.assessments, key=lambda a: (a.skill_id not in crit_ids,
                                                          -a.axes.weighted(job.axis_weights))):
            sd = CATALOG[a.skill_id]
            level = a.axes.weighted(job.axis_weights)
            sources = ", ".join(sorted({e.artifact_label for e in a.evidence}))
            statement = (f"Prouvé par les pièces : {sd.statement['fr']} (voir {sources})." if fr
                         else f"Proven by the material: {sd.statement['en']} (see {sources}).")
            validated.append({"skill_id": a.skill_id, "label": sd.l(job.locale), "statement": statement,
                              "level": level, "level_label": _level_label(level, job.locale), "axes": a.axes,
                              "confidence": a.confidence, "evidence": a.evidence, "source": a.source})

        previous = self.store.get("reports", cand.ref, DashboardReport)
        payload = {
            "engine_version": ENGINE_VERSION, "job_config_version": job.version,
            "compatibility_pct": compat, "skills_component_pct": skills_pct,
            "credentials_component_pct": creds.component_pct, "credentials_weight": creds.weight_applied,
            "confidence": confidence, "evidence_band": band,
            "criteria": [c.model_dump(mode="json") for c in criteria],
            "assessments": [{"skill_id": a.skill_id, "axes": a.axes.model_dump(), "source": a.source,
                             "evidence": [e.model_dump(mode="json", exclude={"excerpt"}) for e in a.evidence]}
                            for a in graph.assessments],
            "escalation": escalation.model_dump(mode="json"),
        }
        entry = self.ledger.append("score", payload, actor=actor, job_id=job.id, candidate_ref=cand.ref)
        report = DashboardReport(
            candidate_ref=cand.ref, job_id=job.id, job_title=job.title, locale=job.locale,
            compatibility_pct=compat, skills_component_pct=skills_pct, credentials_component_pct=creds.component_pct,
            confidence=confidence, evidence_band=band, criteria=criteria, validated_skills=validated,
            gaps=gaps, interview_guide=guide[:3], credentials=creds, warnings=warnings,
            audit=AuditInfo(ledger_entry_id=entry.entry_id, entry_hash=entry.entry_hash, engine_version=ENGINE_VERSION,
                            job_config_version=job.version, escalation=escalation, signals_used=len(l1.signals)),
            decision=previous.decision if previous else None, notice=NOTICE[job.locale],
        )
        self.store.put("reports", cand.ref, report, parent=job.id)
        return report

    # -------------------------------------------------------------------------------------- reads

    def get_candidate(self, ref: str) -> Candidate:
        cand = self.store.get("candidates", ref, Candidate)
        if cand is None:
            raise NotFound(f"candidate {ref} not found")
        return cand

    def get_report(self, ref: str) -> DashboardReport:
        report = self.store.get("reports", ref, DashboardReport)
        if report is None:
            raise NotFound(f"no report for {ref}: run the evaluation first")
        return report

    def get_graph(self, ref: str) -> SkillGraph | None:
        return self.store.get("graphs", ref, SkillGraph)

    def summaries(self, job_id: str) -> list[CandidateSummary]:
        out = []
        for cand in self.store.list("candidates", Candidate, parent=job_id):
            report = self.store.get("reports", cand.ref, DashboardReport)
            pilot = self.latest_pilot_report(cand.ref)
            pilot_index = pilot.get("pilot_index_pct") if pilot else None
            compat = report.compatibility_pct if report else None
            out.append(CandidateSummary(
                candidate_ref=cand.ref,
                compatibility_pct=report.compatibility_pct if report else None,
                confidence=report.confidence if report else None,
                evidence_band=report.evidence_band if report else None,
                top_skills=[s.label for s in report.validated_skills[:4]] if report else [],
                evidence_count=sum(len(s.evidence) for s in report.validated_skills) if report else 0,
                artifacts=len(cand.artifact_ids),
                escalated=bool(report and report.audit.escalation.escalated),
                warnings=len(report.warnings) if report else 0,
                decision=report.decision.decision if report and report.decision else None,
                created_at=cand.created_at,
                pilot_index_pct=pilot_index,
                authenticity_pct=pilot.get("authenticity_pct") if pilot else None,
                verified_pct=round((1 - PILOT_BLEND) * compat + PILOT_BLEND * pilot_index, 1)
                if compat is not None and pilot_index is not None else None,
            ))
        return sorted(out, key=lambda s: (s.compatibility_pct is None, -(s.compatibility_pct or 0)))

    def latest_pilot_report(self, ref: str) -> dict[str, Any] | None:
        rows = self.store.query("SELECT body FROM documents WHERE collection = 'pilot_sessions' AND parent = ?", (ref,))
        reports = [b["report"] for b in (json.loads(r["body"]) for r in rows) if b.get("report")]
        return max(reports, key=lambda r: r.get("generated_at", "")) if reports else None

    def explanation(self, ref: str) -> dict[str, Any]:
        cand = self.get_candidate(ref)
        report = self.store.get("reports", ref, DashboardReport)
        return {
            "candidate_ref": ref, "job_id": cand.job_id,
            "report": report.model_dump(mode="json") if report else None,
            "ledger": [e.model_dump(mode="json") for e in self.ledger.entries(candidate_ref=ref, limit=500)],
            "chain": self.ledger.verify().model_dump(),
        }

    # -------------------------------------------------------------------------------------- human oversight

    def decide(self, ref: str, decision: HumanDecision) -> DashboardReport:
        if len(decision.reviewer.strip()) < 2:
            raise PolicyError("a decision must be signed by a named reviewer")
        report = self.get_report(ref)
        entry = self.ledger.append("human_decision", {
            "decision": decision.decision, "reviewer": decision.reviewer, "rationale": decision.rationale,
            "score_entry": report.audit.ledger_entry_id, "compatibility_pct_seen": report.compatibility_pct,
        }, actor=decision.reviewer, job_id=report.job_id, candidate_ref=ref)
        report.decision = decision.model_copy(update={"ledger_entry_id": entry.entry_id})
        self.store.put("reports", ref, report, parent=report.job_id)
        return report

    def reveal(self, ref: str, reviewer: str, reason: str) -> dict[str, list[str]]:
        report = self.get_report(ref)
        if not report.decision or report.decision.decision not in ("shortlist", "interview"):
            raise PolicyError("identity can only be revealed after a human shortlist or interview decision")
        values = self.vault.reveal(ref)
        self.ledger.append("reidentification", {"reviewer": reviewer, "reason": reason,
                                                "fields": sorted(values)}, actor=reviewer,
                           job_id=report.job_id, candidate_ref=ref)
        return {k: v for k, v in values.items() if k.startswith("IDENTITY_") or k in ("EMAIL", "PHONE", "REPO_URL")}

    def erase(self, ref: str, actor: str) -> dict[str, int]:
        cand = self.get_candidate(ref)
        shredded = self.vault.shred(ref)
        artifacts = self.store.delete_children("artifacts", ref)
        self.store.delete_children("assessments", ref)
        self.store.delete_children("pilot_sessions", ref)
        for media in self.media_dir.glob(f"{ref}-*"):
            media.unlink(missing_ok=True)
        for collection in ("reports", "graphs", "l1", "candidates"):
            self.store.delete(collection, ref)
        self.store.query("DELETE FROM documents WHERE collection = 'explain_tokens' AND body LIKE ?",
                         (f'%"ref": "{ref}"%',))
        self.ledger.append("erasure", {"vault_entries_shredded": shredded, "artifacts_deleted": artifacts,
                                       "legal_basis": "GDPR Art. 17"}, actor=actor, job_id=cand.job_id,
                           candidate_ref=ref)
        return {"vault_entries_shredded": shredded, "artifacts_deleted": artifacts}

    # ---------------------------------------------------------------------------------- GDPR operations

    def purge_expired(self, now: datetime | None = None, actor: str = "retention-policy") -> list[str]:
        """Erase every application whose consented retention period is over (GDPR Art. 5(1)(e))."""
        now = now or utcnow()
        expired = [c.ref for c in self.store.list("candidates", Candidate)
                   if c.consent.recorded_at + timedelta(days=c.consent.retention_days) <= now]
        for ref in expired:
            self.erase(ref, actor)
        if expired:
            self.ledger.append("retention_sweep", {"erased": len(expired), "at": now.isoformat()}, actor=actor)
        return expired

    def export_candidate(self, ref: str, actor: str) -> dict[str, Any]:
        """Everything held about one candidate, for an access or portability request (GDPR Art. 15 and 20)."""
        cand = self.get_candidate(ref)
        report = self.store.get("reports", ref, DashboardReport)
        data = {
            "exported_at": utcnow().isoformat(),
            "candidate_ref": ref,
            "job_id": cand.job_id,
            "consent": cand.consent.model_dump(mode="json"),
            "identity": self.vault.reveal(ref),
            "artifacts": [a.model_dump(mode="json") for a in self.artifacts(ref)],
            "report": report.model_dump(mode="json") if report else None,
            "audit_trail": [e.model_dump(mode="json") for e in self.ledger.entries(candidate_ref=ref, limit=1000)],
            "verification_tests": [json.loads(r["body"]) for r in self.store.query(
                "SELECT body FROM documents WHERE collection = 'assessments' AND parent = ?", (ref,))],
            "ai_pilot_tests": [json.loads(r["body"]) for r in self.store.query(
                "SELECT body FROM documents WHERE collection = 'pilot_sessions' AND parent = ?", (ref,))],
        }
        self.ledger.append("data_export", {"legal_basis": "GDPR Art. 15/20", "sections": sorted(data)},
                           actor=actor, job_id=cand.job_id, candidate_ref=ref)
        return data

    # ---------------------------------------------------------------------------------- explanation links

    def create_explanation_link(self, ref: str, actor: str, days: int = 30) -> dict[str, Any]:
        """A private link the candidate can open to see why they got their score (AI Act Art. 86)."""
        report = self.get_report(ref)
        token = secrets.token_urlsafe(24)
        expires = utcnow() + timedelta(days=days)
        self.store.put_json("explain_tokens", hashlib.sha256(token.encode()).hexdigest(),
                            {"ref": ref, "expires_at": expires.isoformat(), "created_at": utcnow().isoformat(),
                             "created_by": actor})
        self.ledger.append("explanation_link", {"expires_at": expires.isoformat()}, actor=actor,
                           job_id=report.job_id, candidate_ref=ref)
        return {"token": token, "path": f"/explication/{token}", "expires_at": expires.isoformat()}

    def explanation_links(self, ref: str) -> list[dict[str, Any]]:
        self.get_candidate(ref)
        rows = self.store.query("SELECT id, body FROM documents WHERE collection = 'explain_tokens'")
        out = []
        for row in rows:
            body = json.loads(row["body"])
            if body.get("ref") == ref:
                out.append({"id": row["id"][:16], "created_at": body.get("created_at"),
                            "created_by": body.get("created_by"), "expires_at": body["expires_at"],
                            "revoked": bool(body.get("revoked_at")), "revoked_at": body.get("revoked_at"),
                            "revoked_by": body.get("revoked_by"),
                            "active": not body.get("revoked_at")
                            and datetime.fromisoformat(body["expires_at"]) > utcnow()})
        return sorted(out, key=lambda x: x["created_at"] or "", reverse=True)

    def revoke_explanation_link(self, ref: str, link_id: str, actor: str) -> bool:
        rows = self.store.query("SELECT id, body FROM documents WHERE collection = 'explain_tokens' AND id LIKE ?",
                                (f"{link_id}%",))
        rows = [r for r in rows if json.loads(r["body"]).get("ref") == ref and len(link_id) >= 12
                and not json.loads(r["body"]).get("revoked_at")]
        for row in rows:  # kept, flagged: the audit view shows who revoked what and when
            body = json.loads(row["body"]) | {"revoked_at": utcnow().isoformat(), "revoked_by": actor}
            self.store.put_json("explain_tokens", row["id"], body)
        if rows:
            report = self.get_report(ref)
            self.ledger.append("explanation_link_revoked", {"link": link_id}, actor=actor, job_id=report.job_id,
                               candidate_ref=ref)
        return bool(rows)

    def public_explanation(self, token: str) -> dict[str, Any]:
        record = self.store.get_json("explain_tokens", hashlib.sha256(token.encode()).hexdigest())
        if not record or record.get("revoked_at") or datetime.fromisoformat(record["expires_at"]) < utcnow():
            raise NotFound("this explanation link is invalid or has expired")
        ref = record["ref"]
        report = self.get_report(ref)  # NotFound after an erasure: the link dies with the data
        self.ledger.append("explanation_viewed", {}, actor="candidate", job_id=report.job_id, candidate_ref=ref)
        entry = self.ledger.get(report.audit.ledger_entry_id)
        return {
            "job_title": report.job_title, "locale": report.locale, "generated_at": report.generated_at.isoformat(),
            "compatibility_pct": report.compatibility_pct, "skills_component_pct": report.skills_component_pct,
            "credentials_component_pct": report.credentials_component_pct,
            "credentials_weight": report.credentials.weight_applied, "evidence_band": report.evidence_band,
            "criteria": [c.model_dump(mode="json") for c in report.criteria],
            "validated_skills": [s.model_dump(mode="json") for s in report.validated_skills],
            "gaps": [g.model_dump(mode="json") for g in report.gaps],
            "credentials": [c.label for c in report.credentials.items],
            "decision": ({"decision": report.decision.decision, "rationale": report.decision.rationale,
                          "decided_at": report.decision.decided_at.isoformat()} if report.decision else None),
            "notice": report.notice,
            "audit": {"score_entry_hash": report.audit.entry_hash, "recorded_at": entry.created_at if entry else None,
                      "ledger_intact": self.ledger.verify().valid, "engine_version": report.audit.engine_version,
                      "job_config_version": report.audit.job_config_version},
            "expires_at": record["expires_at"],
        }

    # ---------------------------------------------------------------------------------- AI Act operations

    def monitoring(self, tests: list[Any] | None = None) -> dict[str, Any]:
        """Post-market monitoring figures (AI Act Art. 72) and evidence that human oversight is real."""
        jobs_out = []
        overrides = decisions = 0
        bands: dict[str, int] = {}
        for job in self.list_jobs():
            ranked = self.summaries(job.id)
            scores = [s.compatibility_pct for s in ranked if s.compatibility_pct is not None]
            decided = [(i, s) for i, s in enumerate(ranked) if s.decision]
            positive = {"shortlist", "interview"}
            # A human override: someone advanced while a higher-ranked candidate was not (or the reverse).
            for i, s in decided:
                decisions += 1
                higher_not_advanced = any(o.decision not in positive for o in ranked[:i] if o.decision)
                if (s.decision in positive and higher_not_advanced) or (s.decision not in positive and i == 0):
                    overrides += 1
            for s in ranked:
                if s.evidence_band:
                    bands[s.evidence_band] = bands.get(s.evidence_band, 0) + 1
            jobs_out.append({"job_id": job.id, "title": job.title, "candidates": len(ranked), "decisions": len(decided),
                             "score_min": min(scores) if scores else None, "score_max": max(scores) if scores else None,
                             "score_median": sorted(scores)[len(scores) // 2] if scores else None})
        finished = [t for t in (tests or []) if getattr(t, "results", None)]
        risk: dict[str, int] = {}
        for t in finished:
            r = t.results.get("integrity", {}).get("risk", "low")
            risk[r] = risk.get(r, 0) + 1
        spent = sum(self.budget.spent(j.id)["usd"] for j in self.list_jobs())
        incidents = self.ledger.count(kind="incident")
        return {"generated_at": utcnow().isoformat(), "jobs": jobs_out, "evidence_bands": bands,
                "human_decisions": decisions, "decisions_departing_from_ranking": overrides,
                "verification_tests_completed": len(finished), "test_integrity_risk": risk,
                "ai_spend_usd": round(spent, 4), "incidents_recorded": incidents,
                "ledger": self.ledger.verify().model_dump()}

    def record_incident(self, actor: str, severity: str, description: str, affected: list[str]) -> dict[str, Any]:
        """Serious-incident journal (AI Act Art. 73): the reporting deadline is computed from awareness."""
        days = {"death": 10, "widespread": 2, "serious": 15}.get(severity, 15)
        deadline = (utcnow() + timedelta(days=days)).isoformat()
        entry = self.ledger.append("incident", {"severity": severity, "description": description[:4000],
                                                "affected_candidates": affected[:200], "report_deadline": deadline,
                                                "legal_basis": "AI Act Art. 73"}, actor=actor)
        return {"entry_id": entry.entry_id, "severity": severity, "report_to_authority_before": deadline}

    def media_path(self, ref: str, artifact_id: str) -> Path | None:
        matches = sorted(self.media_dir.glob(f"{ref}-{artifact_id}.*"))
        return matches[0] if matches else None

    def verify_ledger(self) -> ChainVerification:
        return self.ledger.verify()

    def ledger_entries(self, **kw: Any) -> list[LedgerEntry]:
        return self.ledger.entries(**kw)

    def usage(self, job_id: str) -> dict[str, Any]:
        job = self.get_job(job_id)
        spent = self.budget.spent(job_id)
        return {"job_id": job_id, "budget_usd": job.funnel.job_budget_usd, **spent,
                "provider": self.provider.name if self.provider else "none",
                "model": self.provider.model if self.provider else ""}


def _merge_credentials(items: list[CredentialItem]) -> list[CredentialItem]:
    """The same diploma named in the CV, on LinkedIn and in its own file counts once, as supported."""
    def key(item: CredentialItem) -> str:
        words = re.findall(r"[^\W\d_]{3,}|\d+", item.label.lower())
        return " ".join(w for w in words if not w.startswith(("school", "person")))[:40]

    merged: dict[str, CredentialItem] = {}
    for item in sorted(items, key=lambda i: not i.supported_by_document):
        k = key(item)
        if k in merged:
            continue
        merged[k] = item
    return list(merged.values())[:20]


def _validate_skills(job: JobProfile) -> None:
    unknown = [c.skill_id for c in job.criteria if c.skill_id not in CATALOG]
    if unknown:
        raise PolicyError(f"unknown skill ids: {', '.join(unknown)}")


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode()).hexdigest()


def _level_label(level: float, locale: str) -> str:
    labels = (("Expert", "Expert"), ("Confirmé", "Advanced"), ("Opérationnel", "Proficient"),
              ("Initié", "Developing"), ("Émergent", "Emerging"))
    thresholds = (3.25, 2.5, 1.75, 1.0, 0.0)
    for (fr, en), t in zip(labels, thresholds, strict=True):
        if level >= t:
            return fr if locale == "fr" else en
    return labels[-1][0 if locale == "fr" else 1]
